import sys, json
import pprint as pp
from pathlib import Path
from elasticsearch import Elasticsearch
import ufodata as ud
import openpyxl 
from datetime import datetime


client = Elasticsearch("http://localhost:9200",  basic_auth=("elastic", 'changeme'))
index_name = "ufo-nuforc-2024"


def curate_hour(hour):
    if hour < 6:
        return 'night'
    elif hour < 12:
        return 'morning'
    elif hour < 19:
        return 'afternoon'
    elif hour < 23:
        return 'evening'
    else:
        return 'evening'


def get_csv_array(file):
    print(file)
    array_out = []
    wb = openpyxl.load_workbook(file)
    ws = wb.active
    print('Total number of rows: '+str(ws.max_row)+'. And total number of columns: '+str(ws.max_column))
    headers = [ws.cell(row=1,column=i).value for i in range(1,ws.max_column+1)]
    print("Headers", headers)
    for line in ws.iter_rows(min_row=2, max_row=ws.max_row+1, min_col=1, max_col=ws.max_column+1, values_only=True):
        #print(line)
        object_out = {}
        for j in range(0,len(line)-1) :
            object_out[headers[j].lower().replace(' ','_')] = line[j]
        array_out.append(object_out)
    return array_out


def curate_location(object_in):
    object_out = object_in
    object_out['location'] = f"{object_in['city']} {object_in['state']} {object_in['country']}" 
    try:
        object_out['calc_geo']=( float(object_in['lon']), float(object_in['lat']))
    except:
         print('Error', object_in['lon'],object_in['lat'])
    return object_out
    
def curate_date(object_in, date_str, time_str):
    object_out = object_in
    if date_str in object_in.keys() and object_in[date_str]:
        date_obj = object_in[date_str]   #datetime.strptime("01/01/0001", '%m/%d/%Y')
        object_out['calc_date'] = date_obj
        object_out['calc_year'] = object_in[date_str].year
        object_out['calc_dayofweek'] = object_in[date_str].weekday()
    if time_str in object_in.keys() and object_in[time_str]:
        object_out['calc_hour'] = object_in[time_str].strftime('%H')
        object_out['calc_moment'] = curate_hour(int(object_out['calc_hour']))
    return object_out

def load_csv(file):
    item_list = get_csv_array(file)
    my_id=0
    file_str=str(file)
    for item in item_list:
        my_id +=1
        pp.pprint(item)
        detailled_item = curate_date(item, 'event_date', 'event_time')
        detailled_item = curate_location(detailled_item)
        detailled_item['calc_src'] = 'NUFORC_ErwanDavid'
        del detailled_item['event_date']
        del detailled_item['event_time']
        pp.pprint(detailled_item)
        #pp.pprint(json.dumps(detailled_item))
        client.index(index=index_name, id=my_id, document=detailled_item)
        print('____')

#ud.create_index(client, index_name)

load_csv(sys.argv[1])
